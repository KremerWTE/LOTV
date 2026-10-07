"""Turn docs/manuals/src/*.md into .docx and .pdf next to them (docs/manuals/).

    python tools/manuals/build.py [name-substring ...]
PDF: Markdown -> HTML -> headless Chromium print. DOCX: python-docx, same content.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from playwright.sync_api import sync_playwright

MANUALS = Path(__file__).resolve().parents[2] / "docs" / "manuals"
SRC = MANUALS / "src"
GREEN = RGBColor(0x2D, 0x4F, 0x1E)

CSS = """
@page { size: Letter; margin: 0.7in; }
body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 11pt; line-height: 1.45; color: #1d2a1a; }
h1 { color:#2d4f1e; font-size: 24pt; border-bottom: 3px solid #2d4f1e; padding-bottom: 6px; }
h2 { color:#2d4f1e; font-size: 16pt; margin-top: 28px; page-break-after: avoid; border-bottom:1px solid #cfd8c8; padding-bottom:3px; }
h3 { font-size: 12.5pt; margin: 18px 0 6px; page-break-after: avoid; }
img { max-width: 100%; border: 1px solid #c9d1c3; border-radius: 4px; margin: 6px 0; page-break-inside: avoid; }
blockquote { margin: 10px 0; padding: 8px 12px; background:#fff8e1; border-left: 4px solid #e0a800; }
code { background:#eef2ea; padding:1px 4px; border-radius:3px; font-size: 10pt; }
ul { padding-left: 22px; } li { margin: 3px 0; }
p { margin: 6px 0; }
"""


def inline_runs(par, text: str):
    """Add text to a docx paragraph honouring **bold**, *italic* and `code`."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            run = par.add_run(tok[2:-2]); run.bold = True
        elif tok.startswith("`"):
            run = par.add_run(tok[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9.5)
        elif tok.startswith("*"):
            run = par.add_run(tok[1:-1]); run.italic = True
        else:
            par.add_run(tok)


def shade(par, hex_fill: str):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def build_docx(md_path: Path, out: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    width = sec.page_width - sec.left_margin - sec.right_margin
    lines = md_path.read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("# "):
            h = doc.add_heading(ln[2:], 0)
        elif ln.startswith("## "):
            h = doc.add_heading(ln[3:], 1)
        elif ln.startswith("### "):
            h = doc.add_heading(ln[4:], 2)
        elif (m := re.match(r"!\[(.*?)\]\((.*?)\)", ln)):
            img = MANUALS / m.group(2)
            if img.exists():
                doc.add_picture(str(img), width=min(width, Inches(6.3)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.paragraphs[-1].paragraph_format.keep_with_next = False
        elif ln.startswith("> "):
            p = doc.add_paragraph(); inline_runs(p, ln[2:]); shade(p, "FFF8E1")
            p.paragraph_format.left_indent = Inches(0.15)
        elif ln.startswith("- [ ] "):
            p = doc.add_paragraph(); inline_runs(p, "☐  " + ln[6:])
        elif ln.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); inline_runs(p, ln[2:])
        else:
            # a paragraph may continue on the next line ("  \n" hard break)
            block = [ln.rstrip()]
            while ln.endswith("  ") and i + 1 < len(lines) and lines[i + 1].strip():
                i += 1; ln = lines[i]; block.append(ln.rstrip())
            p = doc.add_paragraph()
            for k, part in enumerate(block):
                if k:
                    p.add_run().add_break()
                inline_runs(p, part.strip())
        i += 1
    for s in ("Heading 1", "Heading 2", "Title"):
        doc.styles[s].font.color.rgb = GREEN
    for para in doc.paragraphs:  # keep each "Step N" heading with its picture
        if para.style.name == "Heading 2":
            para.paragraph_format.keep_with_next = True
    doc.save(out)


def build_pdf(md_path: Path, out: Path, browser):
    html = markdown.markdown(md_path.read_text(encoding="utf-8"), extensions=["extra", "sane_lists", "nl2br"])
    html = html.replace("[ ]", "☐")
    page = browser.new_page()
    full = f"<!doctype html><meta charset=utf-8><title>{md_path.stem[3:].replace(chr(45), chr(32))}</title><style>{CSS}</style><base href='{MANUALS.as_uri()}/'><body>{html}</body>"
    tmp = MANUALS / "src" / "_tmp.html"
    tmp.write_text(full, encoding="utf-8")
    page.goto(tmp.as_uri())
    page.wait_for_load_state("load")
    page.pdf(path=str(out), format="Letter", print_background=True, margin={"top": "0.6in", "bottom": "0.6in", "left": "0.6in", "right": "0.6in"},
             display_header_footer=True, header_template="<span></span>",
             footer_template="<div style='font-size:8px;width:100%;text-align:center;color:#666'>LOTV — <span class='title'></span> — page <span class='pageNumber'></span> of <span class='totalPages'></span></div>")
    page.close()
    tmp.unlink(missing_ok=True)


def main():
    want = sys.argv[1:]
    files = [f for f in sorted(SRC.glob("*.md")) if not want or any(w in f.name for w in want)]
    with sync_playwright() as p:
        b = p.chromium.launch()
        for f in files:
            build_docx(f, MANUALS / f"{f.stem}.docx")
            build_pdf(f, MANUALS / f"{f.stem}.pdf", b)
            print("built", f.stem, flush=True)
        b.close()


if __name__ == "__main__":
    main()
